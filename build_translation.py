from pathlib import Path
import re,json
from pypdf import PdfReader
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path.cwd(); tmp=root/'tmp/translation'; out=root/'output';out.mkdir(exist_ok=True)
source=PdfReader(r'C:\Users\luluxiu\Downloads\MorphQuad-Morphable Quadrotor for Superhuman.pdf')
refs=[]
for n in range(38,47):
    t=source.pages[n].extract_text().split('Acknowledgments')[0]
    t=re.sub(r'(?m)^\s*(?:References|'+str(n+1)+r')\s*$','',t)
    refs.append(t)
rt='\n'.join(refs)
items=re.split(r'(?m)^\s*(\d{1,2})\.\s+',rt)[1:]
refdict={int(items[i]):items[i+1].strip() for i in range(0,len(items),2)}
assert sorted(refdict)==list(range(1,74)),sorted(refdict)
titles=(root/'reference_titles.txt').read_text(encoding='utf-8').splitlines();assert len(titles)==73

doc=Document(); sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.widow_control=True
for st in doc.styles:
    for border in list(st.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
for n,size in [('Title',20),('Heading 1',16),('Heading 2',13),('Heading 3',11.5)]:
    st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True;st._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.append(f)
doc.core_properties.title='MorphQuad 论文完整中文译文';doc.core_properties.subject='正文 图表 参考文献与补充材料';doc.core_properties.author=''

figs={3:(72,558),8:(72,423),10:(72,635),12:(75,582),18:(109,574),22:(109,573),24:(72,601),31:(109,603)}
used=[]
def picture(line):
    parts=line.split();kind=parts[0][1:];n=int(parts[1]);y0=float(parts[2]);y1=float(parts[3]);x0,x1=(72,540)
    if len(parts)==6:x0,x1=map(float,parts[4:])
    if kind=='FIG':y0,y1=figs[n]
    im=Image.open(tmp/f'source-{n}.png');sx=im.width/612;sy=im.height/792
    crop=im.crop((round(x0*sx),round(y0*sy),round(x1*sx),round(y1*sy)))
    fn=tmp/f'{kind}-{n}-{y0:g}-{y1:g}.png';crop.save(fn)
    width=min(6.3,(x1-x0)/72*0.93) if kind=='EQ' else 6.3
    height=width*crop.height/crop.width
    if kind=='FIG' and height>7.3:width*=7.3/height
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(fn),width=Inches(width))
    used.append({'kind':kind,'source_page':n,'box':[x0,y0,x1,y1]})

def table1():
    heads=['平台','I','II','III','机载感知\n与控制','操作能力','扰动抑制']
    rows=[('固定倾转','','','','','',''),('Yang，2018（1，16，24）','✗','✓','✗','✗','✓','✓'),('Lee，2024（27）','✗','✓','✗','✗','✗','✗'),('Veenstra，2026（28）','✗','✗','✗','✗','✓','✗'),('可变倾转','','','','','',''),('Ryll，2014（31，32）','✗','✗','✓','✗','✗','✗'),('Kamel，2018（34）','✗','✗','✗','✗','✓','✗'),('Ji，2019（35）','✗','✓','✗','✗','✗','✗'),('Bodie，2020（3，41，42）','✗','✓','✗','✗','✓','✓'),('Su，2021（38）','✓','✗','✗','✗','✗','✗'),('Zhao，2023（39，40）','✓','✗','✗','✗','✓','✗'),('SkyGauge，2023（36）*','✗','✗','✓','✓','✓','✓'),('Li，2025（33）','✗','✗','✓','✗','✗','✗'),('Lee，2025（2）','✗','✓','✗','✗','✓','✓'),('Aerix，2025（43，44）*','✗','✗','未知','✓','✗','✓'),('我们的工作','','','','','',''),('MorphEUS，2025（45）','✓','✓','✓','✗','✗','✗'),('MorphQuad（本文）','✓','✓','✓','✓','✓','✓')]
    table=doc.add_table(rows=1,cols=7);table.style='Table Grid';table.autofit=False
    widths=[5.5,1.05,1.05,1.1,2.5,2.4,2.4]
    for c,t,w in zip(table.rows[0].cells,heads,widths):c.text=t;c.width=Cm(w)
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cs=table.add_row().cells
        for c,t,w in zip(cs,vals,widths):c.text=t;c.width=Cm(w)
        no=OxmlElement('w:cantSplit');cs[0]._tc.getparent().get_or_add_trPr().append(no)
    for ri,row in enumerate(table.rows):
        for ci,c in enumerate(row.cells):
            for p in c.paragraphs:
                p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                if ci:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(9);r.font.name='Segoe UI Symbol' if ci else 'Times New Roman';r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');r.bold=ri==0

for line in (root/'translation.md').read_text(encoding='utf-8').splitlines():
    if not line:continue
    if line.startswith('!EQ ') or line.startswith('!FIG '):picture(line)
    elif line=='!TABLE1':table1()
    elif line=='!REFERENCES':
        for n in range(1,74):
            p=doc.add_paragraph(f'[{n}] {titles[n-1]}');p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
            for r in p.runs:r.bold=True;r.font.size=Pt(10)
            raw=refdict[n];raw=re.sub(r'-\s*\n\s*([a-z])',r'\1',raw);raw=re.sub(r'\s+',' ',raw).strip();raw=raw.replace('UA Vs','UAVs').replace('UA V','UAV').replace('Y .','Y.').replace('V .','V.')
            raw=raw.replace('https://www.mdpi. com/','https://www.mdpi.com/').replace('https://dev.realsenseai.com/docs/ intel-','https://dev.realsenseai.com/docs/intel-').replace('camera, accessed:','camera，访问日期：').replace('accessed:','访问日期：')
            for a,b in [('Pf¨andler','Pfändler'),('M. ´A.','M. Á.'),('Mart´ınez','Martínez'),('Mart´ın','Martín'),('L´opez','López'),('Sabl´e','Sablé'),('B ¨ulthoff','Bülthoff'),('Salda ˜na','Saldaña'),('10.15607/ rss','10.15607/rss')]:raw=raw.replace(a,b)
            raw=re.sub(r'\.(?=(?:IEEE|The International|Sensors|Electronics|Acta|Drones|Remote|Science|Autonomous|Applied|Journal|Ergonomics|arXiv))','. ',raw)
            raw=re.sub(r'\bin(?=\d{4}|International|IEEE|Robotics|Proceedings|Internet)', 'in ',raw)
            raw=re.sub(r'(?<=[a-zA-Z])\((?=IEEE|Springer|CRC|202[0-9]|201[0-9])',' (',raw)
            raw=re.sub(r'(?<=[a-zA-Z])(?=\d+\()', ' ',raw)
            p=doc.add_paragraph(raw);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(5)
            for r in p.runs:r.font.size=Pt(9)
    elif line.startswith('# '):doc.add_paragraph(line[2:],'Title')
    elif line.startswith('#### '):doc.add_paragraph(line[5:],'Heading 3')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    else:
        iscap=bool(re.match(r'^图 [1-7] (?:我们提出|（|每套|机动性实验|操作实验|抗扰韧性实验)',line) or line.startswith('视频 1 MorphQuad 概览视频。'));p=doc.add_paragraph(line,'Caption' if iscap else 'Normal')
        if line.startswith('译注：') or line.startswith('图内标注译文：') or line.startswith('视频封面标注译文：'):
            for r in p.runs:r.font.size=Pt(9.5)

for p in doc.paragraphs:
    if p.text=='其中 M 为：':p.paragraph_format.keep_with_next=True
final=out/'MorphQuad_完整中文译文.docx';doc.save(final)
audit={'source_pages':55,'references':len(refdict),'translated_reference_titles':len(titles),'figures':sum(x['kind']=='FIG' for x in used),'formula_images':sum(x['kind']=='EQ' for x in used),'paragraphs':len(doc.paragraphs),'characters':sum(len(p.text) for p in doc.paragraphs),'assets':used}
(tmp/'content_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(str(final));print(json.dumps({k:v for k,v in audit.items() if k!='assets'},ensure_ascii=False))
