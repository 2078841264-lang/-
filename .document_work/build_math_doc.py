"""生成面向代码学习的数学说明；使用原生 OMML 公式。"""
from pathlib import Path
import json, re, hashlib, sys
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'文档';OUT.mkdir(exist_ok=True)
WORK=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.1)
sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.3)
sec.header_distance=Cm(1.0);sec.footer_distance=Cm(1.0)
WIDTH=16.2

def font(style,size,east='宋体',latin='Times New Roman',bold=False):
    s=doc.styles[style];s.font.name=latin;s.font.size=Pt(size);s.font.bold=bold;s.font.italic=False;s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().get_or_add_rFonts();rf.set(qn('w:eastAsia'),east)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        if qn('w:'+key) in rf.attrib:del rf.attrib[qn('w:'+key)]
    s.paragraph_format.space_after=Pt(6)
    s.paragraph_format.line_spacing=1.35
font('Normal',11.5)
font('Title',20,'黑体',bold=True)
font('Subtitle',12,'宋体')
font('Heading 1',15,'黑体',bold=True)
font('Heading 2',12,'黑体',bold=True)
font('Caption',10)
doc.styles.add_style('TOC 1',WD_STYLE_TYPE.PARAGRAPH)
font('TOC 1',11)
doc.styles['TOC 1'].paragraph_format.first_line_indent=Pt(0)
doc.styles['TOC 1'].paragraph_format.space_after=Pt(0)
doc.styles['TOC 1'].paragraph_format.space_before=Pt(0)
doc.styles['TOC 1'].paragraph_format.line_spacing=1.0
for n in ('Heading 1','Heading 2'):
    doc.styles[n].paragraph_format.keep_with_next=True
    doc.styles[n].paragraph_format.space_before=Pt(8)
    doc.styles[n].paragraph_format.space_after=Pt(8)
    doc.styles[n].paragraph_format.first_line_indent=Pt(0)
for n in ('Title','Subtitle'):doc.styles[n].paragraph_format.first_line_indent=Pt(0)
normal=doc.styles['Normal'].paragraph_format
normal.first_line_indent=Pt(23)
normal.widow_control=True

header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=header.add_run('双轴倾转四旋翼闭环仿真的数学原理与代码解读');run.font.size=Pt(9)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
for text,field in [('第 ',None),('', 'PAGE'),(' 页',None)]:
    r=footer.add_run(text);r.font.size=Pt(9)
    if field:
        el=OxmlElement('w:fldSimple');el.set(qn('w:instr'),field);r._r.addnext(el)

def el(tag,*children,**attrs):
    x=OxmlElement('m:'+tag)
    for k,v in attrs.items():x.set(qn('m:'+k),str(v))
    for child in children:
        if isinstance(child,list):
            for y in child:x.append(y)
        else:x.append(child)
    return x

def mr(text,roman=False):
    r=el('r');pr=el('rPr')
    if roman:pr.append(el('sty',val='p'))
    r.append(pr)
    wr=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Cambria Math');rf.set(qn('w:hAnsi'),'Cambria Math');wr.append(rf)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'23');wr.append(sz);r.append(wr)
    t=el('t');t.text=text;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');r.append(t)
    return r

class MathParser:
    """小型结构化公式语法，仅实现本文需要的 OMML 结构，不输出裸 LaTeX。"""
    def __init__(self,s):self.s=s;self.i=0
    def rawgroup(self):
        assert self.s[self.i]=='{',self.s[self.i:]
        self.i+=1;start=self.i;depth=1
        while depth:
            c=self.s[self.i];self.i+=1
            if c=='{':depth+=1
            if c=='}':depth-=1
        return self.s[start:self.i-1]
    def arg(self):
        if self.i<len(self.s) and self.s[self.i]=='{':return MathParser(self.rawgroup()).parse()
        ch=self.s[self.i];self.i+=1;return [mr(ch)]
    def parse(self):
        out=[]
        while self.i<len(self.s):
            c=self.s[self.i]
            if c=='{':node=el('box',el('e',self.arg()))
            elif c=='\\':
                self.i+=1;m=re.match('[A-Za-z]+',self.s[self.i:]);assert m,self.s
                cmd=m.group();self.i+=len(cmd)
                if cmd=='frac':node=el('f',el('num',self.arg()),el('den',self.arg()))
                elif cmd=='sqrt':node=el('rad',el('radPr',el('degHide',val='1')),el('deg'),el('e',self.arg()))
                elif cmd in ('dot','ddot','hat','bar'):
                    node=el('acc',el('accPr',el('chr',val={'dot':'̇','ddot':'̈','hat':'̂','bar':'̅'}[cmd])),el('e',self.arg()))
                elif cmd=='rm':node=mr(self.rawgroup(),True)
                elif cmd=='mat':
                    raw=self.rawgroup();matrix=el('m')
                    for row in raw.split(';'):
                        matrix.append(el('mr',[el('e',MathParser(cell).parse()) for cell in row.split('&')]))
                    node=el('d',el('dPr',el('begChr',val='['),el('endChr',val=']')),el('e',matrix))
                else:raise ValueError(cmd)
            else:
                self.i+=1;node=mr(c, c in '0123456789=+−-×·,.;:()[]| /')
            sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in '_^':
                op=self.s[self.i];self.i+=1;a=self.arg()
                if op=='_':sub=a
                else:sup=a
            if sub is not None and sup is not None:node=el('sSubSup',el('e',node),el('sub',sub),el('sup',sup))
            elif sub is not None:node=el('sSub',el('e',node),el('sub',sub))
            elif sup is not None:node=el('sSup',el('e',node),el('sup',sup))
            out.append(node)
        return out

eq_count=0;page_titles=[]
inline_map={s:s for s in ['a_B','a_N','p_N','v_N','F_B','F_N','F_N,d','R_d','ω_d',
    'q_v','q_w','e_R','e_ω','e_p','e_v','K_R','K_ω','λ_F','τ_m','ξ_k','ξ_{k+1}','ξ_max',
    'p_start','p_target','p_d','v_d','a_d','t_k','x_k','u_k','W_mid','W_d','W_actual','F_B,d',
    'J_j','K_ω,j','K_R,j','d_x','d_y','d_z','n_6']}
inline_map.update({'ξ_max':r'ξ_{\rm{max}}','p_start':r'p_{\rm{start}}','p_target':r'p_{\rm{target}}',
    'F_N,d':'F_{N,d}','F_B,d':'F_{B,d}','K_ω,j':'K_{ω,j}','K_R,j':'K_{R,j}',
    'ė_ω':r'\dot{e}_ω','Ṙ_d':r'\dot{R}_d','ω̇_d':r'\dot{ω}_d'})
inline_pattern=re.compile('('+ '|'.join(re.escape(s) for s in sorted(inline_map,key=len,reverse=True))+r')(?![A-Za-z0-9_])')
def add_prose(para,text):
    pos=0
    for match in inline_pattern.finditer(text):
        para.add_run(text[pos:match.start()]);para._p.append(el('oMath',MathParser(inline_map[match.group()]).parse()));pos=match.end()
    para.add_run(text[pos:])
def p(text,boldlead=False):
    x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if boldlead and '：' in text:
        a,b=text.split('：',1);x.add_run(a+'：').bold=True;add_prose(x,b)
    else:add_prose(x,text)
    return x
def h(text):doc.add_heading(text,level=2)
def page(title):
    if len(page_titles)==1:
        toc_title=doc.add_paragraph('目录',style='Title');toc_title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        toc_title.paragraph_format.page_break_before=True
        toc=doc.add_paragraph();toc.paragraph_format.first_line_indent=Pt(0)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z')
        rr=OxmlElement('w:r');tt=OxmlElement('w:t');tt.text='目录将在 Word 中更新';rr.append(tt);field.append(rr);toc._p.append(field)
    page_titles.append(title)
    heading=doc.add_heading(title,level=1);heading.paragraph_format.page_break_before=True
def eq(s):
    global eq_count
    eq_count+=1
    x=doc.add_paragraph();pf=x.paragraph_format
    pf.first_line_indent=Pt(0);pf.line_spacing=1.1;pf.space_before=Pt(5);pf.space_after=Pt(7)
    pf.keep_together=True
    pf.tab_stops.add_tab_stop(Cm(WIDTH/2),WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
    x.add_run('\t');x._p.append(el('oMath',MathParser(s).parse()));x.add_run(f'\t({eq_count})')
    return eq_count
def code(text):
    x=doc.add_paragraph();x.paragraph_format.first_line_indent=Pt(0)
    x.paragraph_format.line_spacing=1.0;x.paragraph_format.keep_together=True
    x.paragraph_format.space_before=Pt(4);x.paragraph_format.space_after=Pt(7)
    r=x.add_run(text);r.font.name='Consolas';r.font.size=Pt(9)
    r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
def mapping(text):p('代码对应：'+text,True)
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+tag)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
        borders.append(b)
    pr.append(borders)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        for cell,text in zip(t.add_row().cells,vals):cell.text=str(text)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for j,cell in enumerate(row.cells):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margin.append(e)
            tcpr.append(margin)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tcpr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.line_spacing=1.15;para.paragraph_format.space_after=Pt(3)
                for r in para.runs:r.font.size=Pt(10);r.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

# 正文定义与文档生成放在独立内容文件，便于逐页校对。
exec((WORK/'content.py').read_text(encoding='utf-8'))
doc.core_properties.title='双轴倾转四旋翼闭环仿真的数学原理与代码解读'
doc.core_properties.subject='FRD NED 坐标下的位置 PID 与 SO3 控制及推力分配推导'
doc.core_properties.author='技术说明'
# 清除默认模板的标题装饰线，保持纯黑论文式标题。
for root in [doc.styles.element,doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
dest=OUT/'双轴倾转四旋翼数学原理与代码解读.docx'
doc.save(dest)
(WORK/'manifest.json').write_text(json.dumps({'equations':eq_count,'sections':page_titles,'file':str(dest)},ensure_ascii=False,indent=2),encoding='utf-8')
print('Equations:',eq_count,'sections:',len(page_titles));print(dest)
